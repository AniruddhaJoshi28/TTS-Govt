"""Multi-format document converter: any supported file ➜ PDF bytes.

Supported formats:
  PDF          — passthrough (validated)
  DOCX         — python-docx + pymupdf
  TXT / RTF    — plain-text to PDF via pymupdf
  ODT          — zip/xml parser + pymupdf
  Images       — JPG, PNG, TIFF, BMP, WEBP → PDF via Pillow

Every converter returns raw PDF bytes so the rest of the pipeline (extract.py,
corpus.py, ocr.py) works unchanged.
"""

import io
import logging
import mimetypes
import unicodedata
from pathlib import Path

from PIL import Image

logger = logging.getLogger(__name__)

# Map of extension → canonical type used internally.
EXTENSION_MAP: dict[str, str] = {
    ".pdf":  "pdf",
    ".docx": "docx",
    ".doc":  "doc",
    ".txt":  "txt",
    ".rtf":  "rtf",
    ".odt":  "odt",
    ".jpg":  "image",
    ".jpeg": "image",
    ".png":  "image",
    ".tif":  "image",
    ".tiff": "image",
    ".bmp":  "image",
    ".webp": "image",
}

# All extensions the system accepts (for UI / validation).
ALLOWED_EXTENSIONS: set[str] = set(EXTENSION_MAP.keys())

# Human-readable label for error messages.
ALLOWED_LABEL = "PDF, DOCX, DOC, TXT, RTF, ODT, JPG, PNG, TIFF, BMP, WEBP"

# Hard upper limit on uploaded file size (50 MB).
MAX_FILE_BYTES = 50 * 1024 * 1024


class ConversionError(ValueError):
    """Raised when a file cannot be converted to PDF."""


class UnsupportedFormat(ConversionError):
    """The file extension or content is not a supported document type."""


def detect_type(filename: str, content: bytes) -> str:
    """Determine the canonical type from filename and magic bytes."""
    ext = Path(filename).suffix.lower()
    if ext in EXTENSION_MAP:
        return EXTENSION_MAP[ext]

    # Fallback: sniff magic bytes.
    if content[:5] == b"%PDF-":
        return "pdf"
    if content[:4] == b"PK\x03\x04":
        if b"word/" in content[:2000]:
            return "docx"
        if b"opendocument" in content[:2000].lower():
            return "odt"
        return "docx"
    if content[:2] in (b"\xff\xd8", b"\x89P"):  # JPEG, PNG
        return "image"
    if content[:4] in (b"II*\x00", b"MM\x00*"):  # TIFF
        return "image"
    if content[:4] == b"BM\x00\x00" or content[:4].startswith(b"BM"):
        return "image"

    raise UnsupportedFormat(
        f"Unsupported file type '{ext or '(unknown)'}'. "
        f"Accepted formats: {ALLOWED_LABEL}."
    )


def _validate_size(content: bytes) -> None:
    if len(content) > MAX_FILE_BYTES:
        mb = len(content) / (1024 * 1024)
        raise ConversionError(
            f"File is too large ({mb:.1f} MB). Maximum allowed is "
            f"{MAX_FILE_BYTES / (1024 * 1024):.0f} MB."
        )


# -- individual converters ---------------------------------------------------

def _pdf_passthrough(content: bytes) -> bytes:
    """Validate and return PDF as-is."""
    if not content.lstrip()[:5] == b"%PDF-":
        raise ConversionError("File has a .pdf extension but is not a valid PDF.")
    return content


def _image_to_pdf(content: bytes) -> bytes:
    """Convert a single image file to a one-page PDF."""
    try:
        img = Image.open(io.BytesIO(content))
    except Exception as exc:
        raise ConversionError(f"Cannot open image: {exc}") from exc

    if img.mode in ("RGBA", "LA"):
        background = Image.new("RGB", img.size, (255, 255, 255))
        background.paste(img, mask=img.split()[-1])
        img = background
    elif img.mode != "RGB":
        img = img.convert("RGB")

    buf = io.BytesIO()
    img.save(buf, format="PDF", resolution=150)
    return buf.getvalue()


def _txt_to_pdf(content: bytes) -> bytes:
    """Convert plain text to a well-formatted PDF using pymupdf."""
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("latin-1", errors="replace")

    text = text.strip()
    if not text:
        raise ConversionError("The text file is empty.")

    try:
        import pymupdf
    except ImportError:
        raise ConversionError(
            "Text-to-PDF conversion requires pymupdf. "
            "Install it with: pip install pymupdf"
        )

    doc = pymupdf.open()

    font_buffer = None
    for fp in ["C:/Windows/Fonts/Nirmala.ttc", "C:/Windows/Fonts/arial.ttf"]:
        p = Path(fp)
        if p.is_file():
            try:
                font_buffer = pymupdf.Font(fontfile=str(p)).buffer
                break
            except Exception:
                pass

    # Break into lines and wrap long lines for clean page margins
    raw_lines = text.split("\n")
    lines: list[str] = []
    for l in raw_lines:
        l = l.rstrip()
        if len(l) <= 85:
            lines.append(l)
        else:
            words = l.split(" ")
            cur = ""
            for w in words:
                if len(cur) + len(w) + 1 <= 85:
                    cur = (cur + " " + w).strip()
                else:
                    if cur:
                        lines.append(cur)
                    cur = w
            if cur:
                lines.append(cur)

    lines_per_page = 40
    chunks = [lines[i:i + lines_per_page] for i in range(0, len(lines), lines_per_page)]
    if not chunks:
        chunks = [[""]]

    for chunk in chunks:
        page = doc.new_page(width=595, height=842)  # Standard A4 size
        rect = pymupdf.Rect(54, 54, 541, 788)
        page_text = "\n".join(chunk)
        if font_buffer:
            fontname = "indic_font"
            page.insert_font(fontname=fontname, fontbuffer=font_buffer)
            page.insert_textbox(rect, page_text, fontname=fontname, fontsize=11)
        else:
            page.insert_textbox(rect, page_text, fontsize=11)

    return doc.tobytes()


def _rtf_to_pdf(content: bytes) -> bytes:
    """Convert RTF to PDF — extract text content and render as PDF."""
    import re

    try:
        text = content.decode("utf-8", errors="replace")
    except Exception:
        text = content.decode("latin-1", errors="replace")

    # Strip RTF control words and groups.
    text = re.sub(r"\\[a-z]+\d*\s?", " ", text)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\\\*", "", text)
    text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), text)
    text = text.strip()

    if not text:
        raise ConversionError("The RTF file contains no readable text.")

    return _txt_to_pdf(text.encode("utf-8"))


def _docx_to_pdf(content: bytes) -> bytes:
    """Convert DOCX to PDF using python-docx for text extraction + pymupdf for PDF."""
    try:
        from docx import Document
    except ImportError:
        raise ConversionError(
            "DOCX conversion requires python-docx. "
            "Install it with: pip install python-docx"
        )

    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise ConversionError(f"Cannot read DOCX file: {exc}") from exc

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    if not paragraphs:
        raise ConversionError("The DOCX file contains no readable text.")

    return _txt_to_pdf("\n\n".join(paragraphs).encode("utf-8"))


def _odt_to_pdf(content: bytes) -> bytes:
    """Convert ODT to PDF using standard library zipfile and xml parsing."""
    import xml.etree.ElementTree as ET
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml_data = z.read("content.xml")
        root = ET.fromstring(xml_data)
        paragraphs = []
        for elem in root.iter():
            if elem.tag.endswith("}p") or elem.tag.endswith("}h"):
                t = "".join(elem.itertext()).strip()
                if t:
                    paragraphs.append(t)
    except Exception as exc:
        raise ConversionError(f"Cannot read ODT file: {exc}") from exc

    if not paragraphs:
        raise ConversionError("The ODT file contains no readable text.")

    return _txt_to_pdf("\n\n".join(paragraphs).encode("utf-8"))


def _doc_to_pdf(content: bytes) -> bytes:
    """Convert legacy .doc to PDF via best-effort text extraction."""
    if content[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ConversionError(
            "This .doc file does not appear to be a valid legacy Word document. "
            "If it is actually a .docx file, please rename it with a .docx extension."
        )

    import re
    text_runs = []
    try:
        decoded = content.decode("utf-16-le", errors="ignore")
        runs = re.findall(r"[\w\s.,;:!?()\"'-]{10,}", decoded)
        text_runs.extend(runs)
    except Exception:
        pass

    if not text_runs:
        try:
            decoded = content.decode("latin-1", errors="ignore")
            runs = re.findall(r"[\x20-\x7e]{10,}", decoded)
            text_runs.extend(runs)
        except Exception:
            pass

    if not text_runs:
        raise ConversionError(
            "Cannot extract text from this .doc file. "
            "Please convert it to .docx or .pdf using Microsoft Word or LibreOffice."
        )

    text = "\n".join(text_runs)
    return _txt_to_pdf(text.encode("utf-8"))


# -- main entry point --------------------------------------------------------

_CONVERTERS: dict[str, callable] = {
    "pdf":   _pdf_passthrough,
    "docx":  _docx_to_pdf,
    "doc":   _doc_to_pdf,
    "txt":   _txt_to_pdf,
    "rtf":   _rtf_to_pdf,
    "odt":   _odt_to_pdf,
    "image": _image_to_pdf,
}


def to_pdf(filename: str, content: bytes) -> bytes:
    """Convert any supported document format to PDF bytes.

    Parameters
    ----------
    filename : str
        Original filename with extension (used for type detection).
    content : bytes
        Raw file content.

    Returns
    -------
    bytes
        Valid PDF content ready for the extraction pipeline.

    Raises
    ------
    UnsupportedFormat
        If the file type is not supported.
    ConversionError
        If conversion fails (corrupt file, empty content, etc.).
    """
    if not content:
        raise ConversionError("The uploaded file is empty.")

    _validate_size(content)

    file_type = detect_type(filename, content)
    converter = _CONVERTERS.get(file_type)
    if converter is None:
        raise UnsupportedFormat(f"No converter for type '{file_type}'.")

    logger.info("Converting %s (%s, %d bytes) to PDF", filename, file_type, len(content))
    pdf_bytes = converter(content)

    if not pdf_bytes or not pdf_bytes.lstrip()[:5] == b"%PDF-":
        raise ConversionError(
            f"Conversion of {filename} produced invalid output. "
            "Please try converting the file to PDF manually."
        )

    return pdf_bytes


def detect_language(text: str) -> str:
    """Auto-detect language from extracted text using script analysis.

    Returns the ISO 639-1 code: 'mr', 'hi', 'or', or 'en'.
    """
    devanagari = 0
    odia = 0
    latin = 0
    total = 0

    for ch in text[:5000]:
        cp = ord(ch)
        if 0x0900 <= cp <= 0x097F:
            devanagari += 1
            total += 1
        elif 0x0B00 <= cp <= 0x0B7F:
            odia += 1
            total += 1
        elif ch.isascii() and ch.isalpha():
            latin += 1
            total += 1

    if total < 20:
        return "en"

    if odia > devanagari and odia > latin:
        return "or"
    if devanagari > latin:
        hindi_markers = sum(1 for w in ["है", "हैं", "का", "की", "के", "में", "से"]
                           if w in text[:3000])
        marathi_markers = sum(1 for w in ["आहे", "आहेत", "चा", "ची", "चे", "मध्ये", "करणे"]
                              if w in text[:3000])
        return "hi" if hindi_markers > marathi_markers else "mr"
    return "en"
